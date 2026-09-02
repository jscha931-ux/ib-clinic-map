# -*- coding: utf-8 -*-
"""인터베리알파 구매 거래처 -> 지오코딩 -> clinics.json

원천:
  1) Report 2026 V3 / Master  (Index=Actual, M_Code=9174001) -> 구매 C_code
  2) Report 2026 V3 / Ckey                                   -> 거래처명
  3) 동물병원 주소 뽑기.xlsx / C2* 시트 (SAP BP 마스터)       -> 주소
지오코딩은 geocache.json 에 영구 캐시. 신규 주소만 카카오 로컬 API 호출.
"""
import argparse, glob, json, os, re, shutil, sys, tempfile, time, urllib.parse, urllib.request

import openpyxl

DESKTOP = r"C:\Users\Owner\OneDrive - GREENCROSS VETERINARY PRODUCTS CO.,LTD\바탕 화면"
V3_GLOB = os.path.join(DESKTOP, r"MKT\CA 영업부", "2) Report 2026 V3*.xlsx")
BRAND = os.path.join(DESKTOP, r"MKT\CA 제품\인터베리알파\브랜드 페이지")
ADDR_XLSX = os.path.join(BRAND, "동물병원 주소 뽑기.xlsx")
LEGACY_JSON = os.path.join(BRAND, "clinics.json")

HERE = os.path.dirname(os.path.abspath(__file__))
CACHE = os.path.join(HERE, "geocache.json")
OUT = os.path.join(HERE, "docs", "clinics.json")
UNMATCHED = os.path.join(HERE, "unmatched.csv")

IBA_MCODE = 9174001
KR_BOUNDS = (33.0, 39.7, 124.5, 132.1)  # lat_min, lat_max, lng_min, lng_max


def norm_addr(s):
    """캐시 키. 잉여 공백만 눌러 같은 주소가 두 번 지오코딩되지 않게."""
    if not s:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()


def open_locked(path):
    """OneDrive/Excel 로 잠긴 파일은 임시 복사본으로 연다."""
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True), None
    except PermissionError:
        tmp = os.path.join(tempfile.gettempdir(), "ibcm_" + os.path.basename(path))
        shutil.copy2(path, tmp)
        return openpyxl.load_workbook(tmp, read_only=True, data_only=True), tmp


def header_index(ws):
    hdr = next(ws.iter_rows(max_row=1, values_only=True))
    return {h: i for i, h in enumerate(hdr) if h}


def load_buyers():
    """V3 Master 에서 IBA 를 산 적 있는 C_code -> (거래처명, 최종구매일)"""
    src = sorted(glob.glob(V3_GLOB))[-1]
    print("[V3] " + os.path.basename(src))
    wb, tmp = open_locked(src)
    try:
        ws = wb["Master"]
        i = header_index(ws)
        buyers = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[i["Index"]] != "Actual" or r[i["M_Code"]] != IBA_MCODE:
                continue
            code = r[i["C_Code"]]
            if not code:
                continue
            d = r[i["Date"]]
            prev = buyers.get(code)
            if code not in buyers or (d and prev and d > prev) or (d and prev is None):
                buyers[code] = d

        ws = wb["Ckey"]
        i = header_index(ws)
        names = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            c = r[i.get("C_code", 0)]
            if c and c not in names:
                names[c] = r[i["Customer"]]
    finally:
        wb.close()
        if tmp and os.path.exists(tmp):
            os.remove(tmp)
    return {str(c): (names.get(c) or "", d) for c, d in buyers.items()}


def load_addresses():
    """SAP BP 마스터: C_code -> (주소, 전화)"""
    wb, tmp = open_locked(ADDR_XLSX)
    try:
        sheet = next(s for s in wb.sheetnames if s.startswith("C2"))
        print("[주소마스터] " + os.path.basename(ADDR_XLSX) + " / " + sheet)
        ws = wb[sheet]
        out = {}
        for r in ws.iter_rows(min_row=2, max_col=8, values_only=True):
            if not r[0]:
                continue
            addr = r[5] or r[6]  # F(청구처 주소) 우선, 없으면 G(납품처 주소)
            out[str(r[0]).strip()] = (norm_addr(addr), r[7] or "")
    finally:
        wb.close()
        if tmp and os.path.exists(tmp):
            os.remove(tmp)
    return out


def load_cache():
    cache = {}
    if os.path.exists(CACHE):
        cache = json.load(open(CACHE, encoding="utf-8"))
    # 최초 1회: 예전 clinics.json 좌표를 캐시로 흡수 (재지오코딩 방지)
    if not cache and os.path.exists(LEGACY_JSON):
        for c in json.load(open(LEGACY_JSON, encoding="utf-8")):
            a, lat, lng = norm_addr(c.get("address")), c.get("lat"), c.get("lng")
            if a and lat and lng:
                cache[a] = {"lat": float(lat), "lng": float(lng), "src": "legacy"}
        print("[캐시] 기존 clinics.json 에서 %d건 시드" % len(cache))
    return cache


def kakao_geocode(addr, key):
    url = "https://dapi.kakao.com/v2/local/search/address.json?" + urllib.parse.urlencode({"query": addr})
    req = urllib.request.Request(url, headers={"Authorization": "KakaoAK " + key})
    with urllib.request.urlopen(req, timeout=10) as resp:
        docs = json.load(resp).get("documents") or []
    if not docs:
        return None
    d = docs[0]
    lat, lng = float(d["y"]), float(d["x"])
    lo_lat, hi_lat, lo_lng, hi_lng = KR_BOUNDS
    if not (lo_lat <= lat <= hi_lat and lo_lng <= lng <= hi_lng):
        return None  # 한국 밖 = 오지오코딩
    return {"lat": lat, "lng": lng, "src": "kakao"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--key", default=os.environ.get("KAKAO_REST_KEY"), help="카카오 REST API 키")
    ap.add_argument("--dry-run", action="store_true", help="지오코딩 호출 없이 현황만 집계")
    a = ap.parse_args()

    buyers = load_buyers()
    addrs = load_addresses()
    cache = load_cache()

    clinics, no_addr, need = [], [], []
    for code, (name, last) in sorted(buyers.items()):
        addr, phone = addrs.get(code, ("", ""))
        if not addr:
            no_addr.append((code, name))
            continue
        if addr not in cache:
            need.append(addr)
        clinics.append({"code": code, "name": name, "address": addr, "phone": phone})

    print("\n구매 거래처 %d곳 | 주소 있음 %d | 주소 없음 %d | 지오코딩 필요 %d건"
          % (len(buyers), len(clinics), len(no_addr), len(set(need))))

    if need and not a.dry_run and not a.key:
        print("! 카카오 REST API 키가 없어 신규 %d건은 건너뜁니다 (--key 또는 KAKAO_REST_KEY)" % len(set(need)))
        need = []

    if need and not a.dry_run:
        todo = sorted(set(need))
        for n, addr in enumerate(todo, 1):
            try:
                res = kakao_geocode(addr, a.key)
            except Exception as e:
                res = None
                print("  ! %s -> %s" % (addr, e))
            if res:
                cache[addr] = res
            print("  [%d/%d] %s %s" % (n, len(todo), "OK  " if res else "FAIL", addr))
            time.sleep(0.05)
        json.dump(cache, open(CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    out, failed = [], []
    for c in clinics:
        hit = cache.get(c["address"])
        if not hit:
            failed.append(c)
            continue
        out.append({"id": str(len(out) + 1), "name": c["name"], "address": c["address"],
                    "phone": "", "lat": hit["lat"], "lng": hit["lng"]})

    prev = json.load(open(OUT, encoding="utf-8")) if os.path.exists(OUT) else []
    if not a.dry_run:
        os.makedirs(os.path.dirname(OUT), exist_ok=True)
        json.dump(out, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        with open(UNMATCHED, "w", encoding="utf-8-sig") as f:
            f.write("사유,거래처코드,거래처명,주소\n")
            for code, name in no_addr:
                f.write("주소없음,%s,%s,\n" % (code, name))
            for c in failed:
                f.write("지오코딩실패,%s,%s,%s\n" % (c["code"], c["name"], c["address"]))

    print(">> clinics.json %d곳 (이전 %d곳, %+d)" % (len(out), len(prev), len(out) - len(prev)))
    print(">> 보완 필요 %d건 -> %s" % (len(no_addr) + len(failed), os.path.basename(UNMATCHED)))


def _selftest():
    assert norm_addr("  서울시   강남구  1 ") == "서울시 강남구 1"
    assert norm_addr(None) == ""
    lo_lat, hi_lat, lo_lng, hi_lng = KR_BOUNDS
    assert lo_lat <= 37.5 <= hi_lat and lo_lng <= 127.0 <= hi_lng  # 서울 = 통과
    assert not (lo_lat <= 40.7 <= hi_lat)  # 뉴욕 = 차단
    print("selftest ok")


if __name__ == "__main__":
    if "--self-test" in sys.argv:
        _selftest()
    else:
        main()
